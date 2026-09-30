"""S0 intake through the sandbox: PDFs, images, archives, Office files, DXF, HEIC and the limits.

The fuzz/limit acceptance of Phase 2 (zip bomb, path traversal, symlink, oversized image,
malformed PDF → coded errors) is covered here.
"""

from __future__ import annotations

import hashlib
import io
import itertools
import json
import stat
import zipfile
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from archrender.assets import font_path
from archrender.core.config import Settings
from archrender.core.errors import ArchRenderError, ErrorCode
from archrender.core.schemas.document import Word
from archrender.ingest.intake import ingest_file, page_refs, tiles
from archrender.ingest.sandbox import run_tool
from archrender.pipeline.services import Services
from tests.conftest import require_tool
from tests.helpers import MINIMAL_DXF

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


@pytest.fixture
def svc(settings: Settings, project_id: str) -> Services:
    return Services.create(settings)


def _ingest(svc: Services, tmp: Path, name: str, data: bytes):  # type: ignore[no-untyped-def]
    p = tmp / f"upload_{name}"
    p.write_bytes(data)
    return ingest_file(svc, "prj_test", p, name)


def _pdf(
    pages: int = 1, *, encrypt: str | None = None, size: tuple[float, float] | None = None
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("DejaVu", str(font_path("sans"))))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=size or A4, encrypt=encrypt)
    for i in range(pages):
        c.setFont("DejaVu", 20)
        c.drawString(100, 700, f"Yatak Odası {i + 1}")
        c.saveState()
        c.translate(300, 400)
        c.rotate(90)
        c.drawString(0, 0, "3,50")
        c.restoreState()
        c.setLineWidth(3)
        c.line(50, 50, 500, 50)
        c.showPage()
    c.save()
    return buf.getvalue()


def _words(svc: Services, page) -> list[Word]:  # type: ignore[no-untyped-def]
    return [
        Word.model_validate(w) for w in json.loads(svc.store("prj_test").read_bytes(page.words))
    ]


# ---------------------------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------------------------
def test_pdf_pages_rasterised_at_300_dpi_with_text_layer_in_pixels(
    svc: Services, tmp_path: Path
) -> None:
    res = _ingest(svc, tmp_path, "Kat Planı.pdf", _pdf(2))
    assert res.kind == "pdf" and res.pages == 2 and not res.deduplicated
    pages = page_refs(svc, "prj_test", res.document_id)
    p = pages[0]
    # A4 (595.28 × 841.89 pt) at 300 DPI = 2480.3 × 3507.9 px
    assert p.dpi == 300.0 and abs(p.width_px - 2480.3) <= 1 and abs(p.height_px - 3507.9) <= 1  # type: ignore[operator]
    assert p.raster is not None and p.tiles and p.meta["preview"]
    assert abs(p.width_mm - 210.0) < 0.5 and abs(p.height_mm - 297.0) < 0.5  # type: ignore[operator]
    words = {w.text: w for w in _words(svc, p)}
    assert {"Yatak", "Odası", "1", "3,50"} <= words.keys()  # Turkish ı survives
    # "Yatak" starts at x = 100 pt on the baseline y = 700 pt (y up) → x ≈ 417 px, baseline ≈ 592 px
    w = words["Yatak"]
    baseline = (842 - 700) * 300 / 72
    assert abs(w.x0 - 100 * 300 / 72) < 8 and w.y0 < baseline < w.y1  # box spans ascent + descent
    assert w.angle_deg == 0.0 and words["3,50"].angle_deg == 90.0
    vector = json.loads(svc.store("prj_test").read_bytes(p.content))  # type: ignore[arg-type]
    assert vector["objects"]["path"] >= 1 and vector["stroke_width_pt"]["p50"] == 3.0


def test_pdf_dedupe(svc: Services, tmp_path: Path) -> None:
    data = _pdf()
    a = _ingest(svc, tmp_path, "a.pdf", data)
    b = _ingest(svc, tmp_path, "b.pdf", data)
    assert b.deduplicated and b.document_id == a.document_id and b.pages == 1


@pytest.mark.parametrize(
    ("data", "code"),
    [
        (b"%PDF-1.7\n1 0 obj << /Type /Catalog >> garbage", ErrorCode.INGEST_CORRUPT),
        (b"%PDF-1.4\n" + bytes(range(256)) * 40, ErrorCode.INGEST_CORRUPT),
    ],
)
def test_malformed_pdf_is_rejected_with_a_coded_error(
    svc: Services, tmp_path: Path, data: bytes, code: ErrorCode
) -> None:
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "broken.pdf", data)
    assert e.value.code == code and e.value.fix_hint
    assert not svc.db.query("SELECT id FROM documents")  # nothing registered
    assert not svc.store("prj_test").exists(hashlib.sha256(data).hexdigest())  # never stored


def test_encrypted_pdf(svc: Services, tmp_path: Path) -> None:
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "locked.pdf", _pdf(encrypt="gizli"))
    assert e.value.code == ErrorCode.INGEST_ENCRYPTED


def test_pdf_page_limit(svc: Services, tmp_path: Path) -> None:
    svc.settings.pdf_max_pages = 2
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "set.pdf", _pdf(3))
    assert e.value.code == ErrorCode.INGEST_LIMIT_EXCEEDED


def test_oversized_sheet_gets_a_lower_dpi_and_a_warning(svc: Services, tmp_path: Path) -> None:
    svc.settings.pdf_max_page_pixels = 2_000_000
    res = _ingest(svc, tmp_path, "A0.pdf", _pdf(size=(2384, 3370)))  # A0 in points
    p = page_refs(svc, "prj_test", res.document_id)[0]
    assert p.dpi is not None and p.dpi < 300 and p.meta["dpi_reduced"]
    assert p.width_px * p.height_px <= 2_000_000  # type: ignore[operator]
    assert any("DPI" in w for w in res.warnings)


# ---------------------------------------------------------------------------------------------
# images
# ---------------------------------------------------------------------------------------------
def _jpeg_with_orientation(orientation: int) -> bytes:
    img = Image.new("RGB", (60, 30), (255, 0, 0))
    img.paste((0, 0, 255), (0, 0, 10, 30))  # blue stripe on the left
    exif = Image.Exif()
    exif[0x0112] = orientation
    exif[0x010F] = "Apple"
    exif[0x0110] = "iPhone 15"
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif, quality=95)
    return buf.getvalue()


def test_exif_orientation_applied_and_camera_recorded(svc: Services, tmp_path: Path) -> None:
    res = _ingest(svc, tmp_path, "foto.jpg", _jpeg_with_orientation(6))  # rotate 90° CW to display
    p = page_refs(svc, "prj_test", res.document_id)[0]
    assert (p.width_px, p.height_px) == (30, 60)
    raster = Image.open(svc.store("prj_test").path(p.raster))  # type: ignore[arg-type]
    top = np.asarray(raster)[:5].mean(axis=(0, 1))
    assert top[2] > 150 and top[0] < 100  # blue stripe moved to the top
    doc = svc.db.one("SELECT meta_json FROM documents WHERE id = ?", (res.document_id,))
    assert json.loads(doc["meta_json"])["exif"]["model"] == "iPhone 15"


def test_png_alpha_flattened_and_multipage_tiff(svc: Services, tmp_path: Path) -> None:
    buf = io.BytesIO()
    Image.new("RGBA", (8, 8), (0, 0, 0, 0)).save(buf, "PNG")
    res = _ingest(svc, tmp_path, "logo.png", buf.getvalue())
    p = page_refs(svc, "prj_test", res.document_id)[0]
    assert p.meta["alpha_flattened"]
    assert Image.open(svc.store("prj_test").path(p.raster)).getpixel((0, 0)) == (255, 255, 255)  # type: ignore[arg-type]
    tif = io.BytesIO()
    frames = [Image.new("L", (20, 10), v) for v in (0, 128)]
    frames[0].save(tif, "TIFF", save_all=True, append_images=frames[1:])
    res = _ingest(svc, tmp_path, "tarama.tif", tif.getvalue())
    assert res.pages == 2


def test_oversized_image_is_rejected(svc: Services, tmp_path: Path) -> None:
    svc.settings.image_max_pixels = 10_000
    buf = io.BytesIO()
    Image.new("RGB", (400, 400)).save(buf, "PNG")  # 160 000 px > 2 × limit → bomb error
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "huge.png", buf.getvalue())
    assert e.value.code == ErrorCode.INGEST_LIMIT_EXCEEDED


def test_truncated_image_is_rejected(svc: Services, tmp_path: Path) -> None:
    buf = io.BytesIO()
    Image.new("RGB", (200, 200), (10, 20, 30)).save(buf, "PNG")
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "cut.png", buf.getvalue()[:300])
    assert e.value.code == ErrorCode.INGEST_CORRUPT


def test_heic_decoded_by_libheif(svc: Services, tmp_path: Path) -> None:
    require_tool("heif-dec", "heif-convert")
    res = _ingest(svc, tmp_path, "telefon.heic", (FIXTURES / "sample.heic").read_bytes())
    p = page_refs(svc, "prj_test", res.document_id)[0]
    assert (p.width_px, p.height_px) == (96, 64) and p.meta["converted_from"] == "heic"
    px = Image.open(svc.store("prj_test").path(p.raster)).getpixel((25, 30))  # type: ignore[arg-type]
    assert px[2] > 150 and px[0] < 80  # the blue rectangle


# ---------------------------------------------------------------------------------------------
# archives
# ---------------------------------------------------------------------------------------------
def _zip(entries: list[tuple[zipfile.ZipInfo | str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in entries:
            zf.writestr(info, data)
    return buf.getvalue()


def test_zip_unpacked_recursively_with_skips_reported(svc: Services, tmp_path: Path) -> None:
    inner = _zip([("içerik/plan.dxf", MINIMAL_DXF)])
    png = io.BytesIO()
    Image.new("RGB", (4, 4)).save(png, "PNG")
    data = _zip(
        [
            ("proje/Kat Planı.pdf", _pdf()),
            ("proje/görsel.png", png.getvalue()),
            ("proje/ek.zip", inner),
            ("proje/setup.exe", b"MZ\x90\x00" + bytes(64)),
            ("__MACOSX/._x", b"junk"),
        ]
    )
    res = _ingest(svc, tmp_path, "teslim.zip", data)
    kinds = sorted(
        r["kind"] for r in svc.db.query("SELECT kind FROM documents WHERE parent_id IS NOT NULL")
    )
    assert kinds == ["dxf", "pdf", "png", "zip"]
    assert len(res.children) == 3  # pdf, png, zip (the dxf is a grandchild)
    assert [s.name for s in res.skipped] == ["proje/setup.exe"]
    assert res.skipped[0].code == ErrorCode.INGEST_UNSUPPORTED_TYPE.value
    names = {r["filename"] for r in svc.db.query("SELECT filename FROM documents")}
    assert {"Kat Planı.pdf", "görsel.png", "plan.dxf"} <= names


def _info(name: str, mode: int | None = None) -> zipfile.ZipInfo:
    zi = zipfile.ZipInfo(name)
    if mode is not None:
        zi.external_attr = mode << 16
    return zi


@pytest.mark.parametrize(
    ("entries", "code"),
    [
        ([("../../etc/evil.txt", b"x")], ErrorCode.INGEST_UNSAFE_ARCHIVE),
        ([("/abs/path.txt", b"x")], ErrorCode.INGEST_UNSAFE_ARCHIVE),
        ([(_info("link", stat.S_IFLNK | 0o777), b"/etc/passwd")], ErrorCode.INGEST_UNSAFE_ARCHIVE),
        ([(_info("dev", stat.S_IFCHR | 0o600), b"")], ErrorCode.INGEST_UNSAFE_ARCHIVE),
        ([("bomb.bin", bytes(30 * 1024 * 1024))], ErrorCode.INGEST_UNSAFE_ARCHIVE),
    ],
    ids=["traversal", "absolute", "symlink", "device", "bomb"],
)
def test_unsafe_archives_are_rejected(
    svc: Services,
    tmp_path: Path,
    entries: list[tuple[zipfile.ZipInfo | str, bytes]],
    code: ErrorCode,
) -> None:
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "x.zip", _zip(entries))
    assert e.value.code == code
    assert not svc.db.query("SELECT id FROM documents")


def test_encrypted_entry_and_entry_limit(svc: Services, tmp_path: Path) -> None:
    # zipfile never writes encrypted entries, so set the "encrypted" flag bit in the headers
    raw = bytearray(_zip([("secret.txt", b"data")]))
    for sig, off in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        i = raw.index(sig) + off
        raw[i] |= 0x1
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "enc.zip", bytes(raw))
    assert e.value.code == ErrorCode.INGEST_ENCRYPTED
    svc.settings.zip_max_entries = 3
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "many.zip", _zip([(f"f{i}.txt", b"x") for i in range(5)]))
    assert e.value.code == ErrorCode.INGEST_UNSAFE_ARCHIVE


def test_nesting_deeper_than_the_limit_is_skipped(svc: Services, tmp_path: Path) -> None:
    data = _zip([("deep.txt", b"en derin")])
    for level in range(4):
        data = _zip([(f"level{level}.zip", data)])
    res = _ingest(svc, tmp_path, "nested.zip", data)
    assert res.skipped and res.skipped[0].code == ErrorCode.INGEST_UNSAFE_ARCHIVE.value
    assert "nested deeper" in res.skipped[0].message


# ---------------------------------------------------------------------------------------------
# Office, DXF, text
# ---------------------------------------------------------------------------------------------
def test_docx_with_table_and_embedded_image(svc: Services, tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Proje özeti: salon aydınlık, meşe parke.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Mahal", "Zemin"
    t.cell(1, 0).text, t.cell(1, 1).text = "Salon", "Meşe parke"
    png = tmp_path / "moodboard.png"
    Image.new("RGB", (32, 32), (120, 90, 60)).save(png)
    d.add_picture(str(png))
    buf = io.BytesIO()
    d.save(buf)
    res = _ingest(svc, tmp_path, "Tasarım Özeti.docx", buf.getvalue())
    content = json.loads(
        svc.store("prj_test").read_bytes(page_refs(svc, "prj_test", res.document_id)[0].content)
    )  # type: ignore[arg-type]
    assert content["paragraphs"][0]["text"].startswith("Proje özeti")
    assert content["tables"][0][1] == ["Salon", "Meşe parke"]
    assert len(res.children) == 1
    child = svc.db.one("SELECT kind, meta_json FROM documents WHERE id = ?", (res.children[0],))
    assert (
        child["kind"] == "png"
        and json.loads(child["meta_json"])["embedded_in"] == "Tasarım Özeti.docx"
    )


def test_xlsx_sheets_and_xlsm_read_as_data(svc: Services, tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Malzeme Listesi"
    ws.append(["Mahal", "Zemin", "Duvar"])
    ws.append(["Salon", "Meşe parke", "Boya RAL 9010"])
    buf = io.BytesIO()
    wb.save(buf)
    res = _ingest(svc, tmp_path, "malzeme.xlsx", buf.getvalue())
    page = page_refs(svc, "prj_test", res.document_id)[0]
    rows = json.loads(svc.store("prj_test").read_bytes(page.content))["rows"]  # type: ignore[arg-type]
    assert page.kind == "sheet" and rows[1] == ["Salon", "Meşe parke", "Boya RAL 9010"]
    # the same workbook with a VBA project: accepted as data, VBA ignored
    xlsm = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(xlsm, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item))
        dst.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0 fake vba")
    res = _ingest(svc, tmp_path, "makro.xlsm", xlsm.getvalue())
    doc = svc.db.one("SELECT kind, meta_json FROM documents WHERE id = ?", (res.document_id,))
    assert doc["kind"] == "xlsm" and json.loads(doc["meta_json"])["vba_ignored"] is True


def test_docm_is_rejected(svc: Services, tmp_path: Path) -> None:
    import docx

    buf = io.BytesIO()
    docx.Document().save(buf)
    docm = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(docm, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item))
        dst.writestr("word/vbaProject.bin", b"vba")
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "rapor.docm", docm.getvalue())
    assert e.value.code == ErrorCode.INGEST_UNSUPPORTED_TYPE and ".docx" in e.value.fix_hint


def test_pptx_slides_and_pictures(svc: Services, tmp_path: Path) -> None:
    import pptx
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Konsept: Sıcak minimalizm"
    png = tmp_path / "mb.png"
    Image.new("RGB", (16, 16), (200, 180, 150)).save(png)
    slide.shapes.add_picture(str(png), Inches(1), Inches(1))
    buf = io.BytesIO()
    prs.save(buf)
    res = _ingest(svc, tmp_path, "konsept.pptx", buf.getvalue())
    page = page_refs(svc, "prj_test", res.document_id)[0]
    content = json.loads(svc.store("prj_test").read_bytes(page.content))  # type: ignore[arg-type]
    assert content["texts"] == ["Konsept: Sıcak minimalizm"] and page.meta["pictures"] == 1
    assert len(res.children) == 1


def test_dxf_summary_texts_attribs_and_preview(svc: Services, tmp_path: Path) -> None:
    import ezdxf

    doc = ezdxf.new("R2010", units=6)
    msp = doc.modelspace()
    doc.layers.add("DUVAR")
    msp.add_lwpolyline([(0, 0), (5, 0), (5, 4), (0, 4)], close=True, dxfattribs={"layer": "DUVAR"})
    msp.add_text("SALON", dxfattribs={"height": 0.3, "layer": "YAZI"}).set_placement((2, 2))
    blk = doc.blocks.new("KAPI")
    blk.add_line((0, 0), (0.9, 0))
    blk.add_attdef("NO", (0, 0.2), dxfattribs={"height": 0.15})
    msp.add_blockref("KAPI", (1, 0)).add_auto_attribs({"NO": "K1"})
    path = tmp_path / "plan.dxf"
    doc.saveas(path)
    res = _ingest(svc, tmp_path, "Zemin Kat.dxf", path.read_bytes())
    page = page_refs(svc, "prj_test", res.document_id)[0]
    summary = json.loads(svc.store("prj_test").read_bytes(page.content))  # type: ignore[arg-type]
    assert summary["insunits"] == 6 and summary["unit_mm"] == 1000.0
    assert summary["layers"]["DUVAR"] == {"LWPOLYLINE": 1}
    texts = {t["text"]: t for t in summary["texts"]}
    assert texts["SALON"]["layer"] == "YAZI" and texts["K1"]["tag"] == "NO"
    assert page.raster is not None and page.width_px and page.width_px > 1000
    img = np.asarray(Image.open(svc.store("prj_test").path(page.raster)))
    assert (img < 128).any()  # the polyline was drawn


def test_dwg_converted_by_libredwg(svc: Services, tmp_path: Path) -> None:
    require_tool("dwg2dxf", "/opt/libredwg/bin/dwg2dxf")
    res = _ingest(svc, tmp_path, "Kat Planı.dwg", (FIXTURES / "sample.dwg").read_bytes())
    assert res.kind == "dwg" and res.pages == 1
    page = page_refs(svc, "prj_test", res.document_id)[0]
    assert page.kind == "dxf" and page.meta["converted_by"].startswith("LibreDWG 0.14")
    summary = json.loads(svc.store("prj_test").read_bytes(page.content))  # type: ignore[arg-type]
    assert summary["layers"]["DUVAR"] == {"LINE": 4} and summary["texts"][0]["text"] == "SALON"
    assert svc.store("prj_test").exists(page.meta["dxf"]["sha256"])  # converted DXF kept for S2
    # LibreDWG 0.14 writes handle 0 for the R2000 ENDBLK records; intake repairs and reports it
    assert summary["repaired_handles"] >= 1 and any("handle 0" in w for w in res.warnings)


def test_dwg_without_libredwg_gives_an_actionable_error(
    svc: Services, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import archrender.ingest.intake as intake

    monkeypatch.setattr(intake, "which", lambda name: None)
    with pytest.raises(ArchRenderError) as e:
        _ingest(svc, tmp_path, "plan.dwg", (FIXTURES / "sample.dwg").read_bytes())
    assert e.value.code == ErrorCode.INGEST_TOOL_MISSING and "DXF or PDF" in e.value.fix_hint


def test_minimal_dxf_and_plain_text(svc: Services, tmp_path: Path) -> None:
    res = _ingest(svc, tmp_path, "k.dxf", MINIMAL_DXF)
    assert res.kind == "dxf" and res.pages == 1
    res = _ingest(svc, tmp_path, "notlar.txt", b"Mutfak: beyaz dolaplar\n")
    page = page_refs(svc, "prj_test", res.document_id)[0]
    assert json.loads(svc.store("prj_test").read_bytes(page.content))["text"].startswith("Mutfak")  # type: ignore[arg-type]


# ---------------------------------------------------------------------------------------------
# sandbox limits and tiling
# ---------------------------------------------------------------------------------------------
def test_sandbox_cpu_limit_and_timeout(settings: Settings, tmp_path: Path) -> None:
    import sys

    settings.sandbox_cpu_s = 1
    with pytest.raises(ArchRenderError) as e:
        run_tool(
            [sys.executable, "-c", "while True: pass"],
            settings,
            tmp_path,
            what="busy loop",
            fix_hint="",
        )
    assert e.value.code == ErrorCode.INGEST_LIMIT_EXCEEDED and "CPU" in e.value.message
    settings.sandbox_cpu_s = 60
    settings.sandbox_timeout_s = 0.5
    with pytest.raises(ArchRenderError) as e:
        run_tool(["sleep", "5"], settings, tmp_path, what="sleep", fix_hint="")
    assert e.value.code == ErrorCode.INGEST_LIMIT_EXCEEDED


def test_sandbox_memory_limit(settings: Settings, tmp_path: Path) -> None:
    import sys

    settings.sandbox_memory_mb = 256
    proc = run_tool(
        [sys.executable, "-c", "x = bytearray(1024 * 1024 * 1024)"],
        settings,
        tmp_path,
        what="alloc",
        fix_hint="",
    )
    assert proc.returncode != 0 and "MemoryError" in proc.stderr


def test_tiles_cover_the_raster_with_overlap() -> None:
    ts = tiles(4000, 2000, 1536, 0.2)
    covered = np.zeros((2000, 4000), bool)
    for t in ts:
        covered[t.y : t.y + t.h, t.x : t.x + t.w] = True
        assert t.x + t.w <= 4000 and t.y + t.h <= 2000
    assert covered.all()
    xs = sorted({t.x for t in ts})
    assert all(b - a <= int(1536 * 0.8) for a, b in itertools.pairwise(xs))
    assert (
        tiles(800, 600, 1536, 0.2) == [tiles(800, 600, 1536, 0.2)[0]]
        and tiles(800, 600, 1536, 0.2)[0].w == 800
    )
