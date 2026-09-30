"""Mutation fuzzing of S0 intake (Phase-2 acceptance: malformed input is rejected with coded errors).

Valid files of every supported kind are truncated, bit-flipped or spliced with random bytes
(seeded, reproducible). Each mutant must either be ingested or be rejected with an
``ArchRenderError`` carrying an ``INGEST_*`` code and a fix hint: never an uncaught exception, a
crash of the worker process or a hang (the sandbox's CPU/wall limits bound each parse).
The unmutated files must also ingest from an upload staging path (``<id>.bin``, no extension).
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from PIL import Image

from archrender.core.config import Settings
from archrender.core.errors import ArchRenderError
from archrender.ingest.intake import ingest_file
from archrender.pipeline.services import Services
from archrender.synth.layout import random_layout
from archrender.synth.sheets import floor_plan_page
from tests.conftest import require_tool
from tests.helpers import MINIMAL_DXF

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
MUTANTS_PER_KIND = 6


def _png() -> bytes:
    buf = io.BytesIO()
    Image.fromarray(np.random.default_rng(1).integers(0, 255, (120, 160, 3), np.uint8)).save(
        buf, "PNG"
    )
    return buf.getvalue()


def _jpeg() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (200, 150), (180, 120, 90)).save(buf, "JPEG", quality=80)
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Tasarım özeti: salon zemini meşe parke.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "KOD", "GENİŞLİK"
    t.cell(1, 0).text, t.cell(1, 1).text = "K1", "90"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["KOD", "GENİŞLİK", "YÜKSEKLİK"])
    ws.append(["K1", 90, 210])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("plan/image.png", _png())
        z.writestr("notes.txt", "Kat yüksekliği 2,80 m")
    return buf.getvalue()


BASES: dict[str, Callable[[], bytes]] = {
    "plan.pdf": lambda: (
        floor_plan_page(np.random.default_rng(5), random_layout(np.random.default_rng(5))).pdf
    ),
    "image.png": _png,
    "photo.jpg": _jpeg,
    "plan.dxf": lambda: (FIXTURES / "sample.dxf").read_bytes(),
    "minimal.dxf": lambda: MINIMAL_DXF,
    "brief.docx": _docx,
    "list.xlsx": _xlsx,
    "bundle.zip": _zip,
}


def _mutants(data: bytes, rng: np.random.Generator) -> list[tuple[str, bytes]]:
    out = []
    for i in range(MUTANTS_PER_KIND):
        kind = ("truncate", "flip", "splice")[i % 3]
        b = bytearray(data)
        if kind == "truncate":
            b = b[: int(rng.integers(1, len(b)))]
        elif kind == "flip":
            for pos in rng.integers(0, len(b), size=max(1, len(b) // 500)):
                b[pos] ^= 1 << int(rng.integers(0, 8))
        else:
            pos = int(rng.integers(0, len(b)))
            b[pos:pos] = rng.integers(
                0, 256, size=int(rng.integers(16, 512)), dtype=np.uint8
            ).tobytes()
        out.append((kind, bytes(b)))
    return out


@pytest.mark.parametrize("name", sorted(BASES))
def test_mutated_files_are_ingested_or_rejected_with_a_coded_error(
    settings: Settings, project_id: str, tmp_path: Path, name: str
) -> None:
    svc = Services.create(settings)
    rng = np.random.default_rng(sorted(BASES).index(name))  # reproducible per kind
    outcomes: list[str] = []
    for i, (kind, data) in enumerate(_mutants(BASES[name](), rng)):
        path = tmp_path / f"m{i}_{name}"
        path.write_bytes(data)
        error: ArchRenderError | None = None
        try:
            ingest_file(svc, project_id, path, f"m{i}_{name}")
        except ArchRenderError as e:
            error = e
        if error is None:
            outcomes.append(f"{kind}:ok")
            continue
        assert error.code.value.startswith("INGEST_"), (kind, error.code, error.message)
        assert error.fix_hint, (kind, error.code, error.message)
        outcomes.append(f"{kind}:{error.code.value}")
    assert len(outcomes) == MUTANTS_PER_KIND, outcomes


@pytest.mark.parametrize("name", sorted(BASES))
def test_unmutated_files_ingest_from_an_upload_staging_path(
    settings: Settings, project_id: str, tmp_path: Path, name: str
) -> None:
    """Uploads are staged as ``<id>.bin``: no parser may depend on the file extension (openpyxl
    did, so every uploaded workbook failed)."""
    svc = Services.create(settings)
    staged = tmp_path / "upl_0123456789abcdef.bin"
    staged.write_bytes(BASES[name]())
    res = ingest_file(svc, project_id, staged, name)
    assert res.pages >= 1 or res.children, res


@pytest.mark.parametrize(
    ("fixture", "tools"),
    [
        ("sample.heic", ("heif-dec", "heif-convert")),
        ("sample.dwg", ("dwg2dxf", "/opt/libredwg/bin/dwg2dxf")),
    ],
)
def test_converted_formats_ingest_from_an_upload_staging_path(
    settings: Settings, project_id: str, tmp_path: Path, fixture: str, tools: tuple[str, ...]
) -> None:
    require_tool(*tools)
    svc = Services.create(settings)
    staged = tmp_path / "upl_fedcba9876543210.bin"
    staged.write_bytes((FIXTURES / fixture).read_bytes())
    assert ingest_file(svc, project_id, staged, fixture).pages == 1
