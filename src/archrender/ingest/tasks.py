"""Parsers for untrusted input. Run only inside the sandbox (``archrender.ingest.sandbox``).

    python -m archrender.ingest.tasks <task>   # JSON request on stdin, JSON reply on stdout

Each task reads one input file, writes its outputs (PNG rasters, JSON) into ``request["out"]`` and
returns a dict::

    {"pages": [{"index", "kind", "width_px", "height_px", "dpi", "width_mm", "height_mm",
                "raster", "preview", "words", "content", "meta"}],
     "children": [{"path", "name"}], "skipped": [{"name", "code", "message", "fix_hint"}],
     "warnings": [...], "meta": {...}}

File names in the reply are relative to ``out``. Coordinates of words are raster pixels (origin
top-left, y down). Nothing here touches the database or the CAS.
"""

from __future__ import annotations

import ctypes
import io
import json
import math
import stat
import sys
import traceback
import unicodedata
import warnings
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path, PurePosixPath
from typing import Any

import numpy as np

from archrender.core.errors import ArchRenderError, ErrorCode

Reply = dict[str, Any]
PREVIEW_PX = 1024
MM_PER_PT = 25.4 / 72.0


# ---------------------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------------------
def _write_json(out: Path, name: str, data: Any) -> str:
    (out / name).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return name


def _save_png(img: Any, out: Path, name: str) -> str:
    img.save(out / name, format="PNG", compress_level=3)
    return name


def _preview(img: Any, out: Path, name: str) -> str:
    from PIL import Image

    p = img.copy()
    p.thumbnail((PREVIEW_PX, PREVIEW_PX), Image.Resampling.LANCZOS)
    return _save_png(p, out, name)


def _empty_reply() -> Reply:
    return {"pages": [], "children": [], "skipped": [], "warnings": [], "meta": {}}


# ---------------------------------------------------------------------------------------------
# PDF (pypdfium2)
# ---------------------------------------------------------------------------------------------
_BREAK_CHARS = {" ", "\t", "\r", "\n", "\x00", " ", "￾", "\x02"}


def _page_to_px(
    page: Any, width: int, height: int
) -> Callable[[float, float], tuple[float, float]]:
    """Affine map PDF user space → raster pixels, honouring /Rotate and the crop box (via PDFium)."""
    import pypdfium2.raw as r

    def dev(x: float, y: float) -> tuple[int, int]:
        dx, dy = ctypes.c_int(), ctypes.c_int()
        r.FPDF_PageToDevice(page.raw, 0, 0, width, height, 0, x, y, dx, dy)
        return dx.value, dy.value

    left, bottom, right, top = page.get_cropbox()
    p0 = np.array(dev(left, bottom), dtype=np.float64)
    px = np.array(dev(right, bottom), dtype=np.float64)
    py = np.array(dev(left, top), dtype=np.float64)
    ex = (px - p0) / max(right - left, 1e-9)
    ey = (py - p0) / max(top - bottom, 1e-9)

    def f(x: float, y: float) -> tuple[float, float]:
        v = p0 + ex * (x - left) + ey * (y - bottom)
        return float(v[0]), float(v[1])

    return f


def _pdf_words(
    page: Any, textpage: Any, to_px: Callable[[float, float], tuple[float, float]]
) -> list[dict[str, Any]]:
    import pypdfium2.raw as r

    n = textpage.count_chars()
    words: list[dict[str, Any]] = []
    cur: list[tuple[str, float, float, float, float, float]] = []

    def flush() -> None:
        if not cur:
            return
        text = "".join(c[0] for c in cur)
        xs = [c[1] for c in cur] + [c[3] for c in cur]
        ys = [c[2] for c in cur] + [c[4] for c in cur]
        angle = cur[0][5]
        words.append(
            {
                "text": unicodedata.normalize("NFC", text),
                "x0": min(xs),
                "y0": min(ys),
                "x1": max(xs),
                "y1": max(ys),
                "angle_deg": angle,
                "confidence": 1.0,
                "source": "pdf_text",
            }
        )
        cur.clear()

    matrix = r.FS_MATRIX()
    ox, oy = to_px(0.0, 0.0)
    for i in range(n):
        ch = textpage.get_text_range(i, 1)
        if not ch or ch in _BREAK_CHARS or r.FPDFText_IsGenerated(textpage.raw, i) == 1:
            flush()
            continue
        left, bottom, right, top = textpage.get_charbox(i, loose=True)
        corners = [to_px(left, bottom), to_px(right, bottom), to_px(right, top), to_px(left, top)]
        cx = [c[0] for c in corners]
        cy = [c[1] for c in corners]
        angle = 0.0
        if r.FPDFText_GetMatrix(textpage.raw, i, matrix):
            ux, uy = to_px(matrix.a, matrix.b)
            angle = round(math.degrees(math.atan2(-(uy - oy), ux - ox)), 1) % 360.0
        box = (ch, min(cx), min(cy), max(cx), max(cy), angle)
        if cur:
            prev = cur[-1]
            h = max(prev[4] - prev[2], prev[3] - prev[1], 1.0)
            vertical = 45.0 <= prev[5] % 180.0 <= 135.0
            if abs(box[5] - prev[5]) > 2.0:
                flush()
            elif vertical:
                gap = min(abs(box[2] - prev[4]), abs(prev[2] - box[4]))
                if abs((box[1] + box[3]) / 2 - (prev[1] + prev[3]) / 2) > 0.6 * h or gap > 0.6 * (
                    prev[3] - prev[1] + 1
                ):
                    flush()
            else:
                gap = box[1] - prev[3]
                if (
                    abs((box[2] + box[4]) / 2 - (prev[2] + prev[4]) / 2) > 0.6 * h
                    or gap > 0.35 * h
                    or gap < -h
                ):
                    flush()
        cur.append(box)
    flush()
    return words


def _pdf_vector_stats(page: Any) -> dict[str, Any]:
    import pypdfium2.raw as r

    counts = {"path": 0, "text": 0, "image": 0, "shading": 0, "form": 0}
    names = {
        r.FPDF_PAGEOBJ_PATH: "path",
        r.FPDF_PAGEOBJ_TEXT: "text",
        r.FPDF_PAGEOBJ_IMAGE: "image",
        r.FPDF_PAGEOBJ_SHADING: "shading",
        r.FPDF_PAGEOBJ_FORM: "form",
    }
    segments = 0
    widths: list[float] = []
    image_area = 0.0
    w, h = page.get_size()
    width = ctypes.c_float()
    for obj in page.get_objects(max_depth=8):
        kind = names.get(obj.type, "other")
        counts[kind] = counts.get(kind, 0) + 1
        if kind == "path":
            segments += max(0, r.FPDFPath_CountSegments(obj.raw))
            if len(widths) < 200_000 and r.FPDFPageObj_GetStrokeWidth(obj.raw, width):
                widths.append(float(width.value))
        elif kind == "image":
            x0, y0, x1, y1 = obj.get_bounds()
            image_area += max(0.0, x1 - x0) * max(0.0, y1 - y0)
    pct = {f"p{q}": float(np.percentile(widths, q)) for q in (50, 90, 99)} if widths else {}
    return {
        "objects": counts,
        "path_segments": segments,
        "stroke_width_pt": pct,
        "image_area_fraction": min(1.0, image_area / max(w * h, 1e-9)),
    }


def task_pdf(req: dict[str, Any]) -> Reply:
    import pypdfium2 as pdfium

    src, out = Path(req["path"]), Path(req["out"])
    try:
        pdf = pdfium.PdfDocument(src)
    except pdfium.PdfiumError as e:
        msg = str(e)
        if "password" in msg.lower():
            raise ArchRenderError(
                ErrorCode.INGEST_ENCRYPTED,
                "The PDF is password-protected.",
                "Remove the password (print to an unprotected PDF) and upload it again.",
            ) from e
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The PDF cannot be read ({msg}).",
            "Re-export or re-print the PDF from the authoring application.",
        ) from e
    n = len(pdf)
    if n == 0:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT, "The PDF has no pages.", "Upload the full drawing set."
        )
    if n > int(req["max_pages"]):
        raise ArchRenderError(
            ErrorCode.INGEST_LIMIT_EXCEEDED,
            f"The PDF has {n} pages; the limit is {req['max_pages']}.",
            "Split the set into several PDFs.",
        )
    reply = _empty_reply()
    reply["meta"] = {"page_count": n}
    for i in range(n):
        page = pdf[i]
        w_pt, h_pt = page.get_size()
        area_in2 = (w_pt / 72.0) * (h_pt / 72.0)
        dpi = float(req["dpi"])
        max_px = float(req["max_page_pixels"])
        reduced = False
        if area_in2 * dpi * dpi > max_px:
            dpi = math.floor(math.sqrt(max_px / area_in2))
            reduced = True
        bitmap = page.render(scale=dpi / 72.0, may_draw_forms=True)
        img = bitmap.to_pil().convert("RGB")
        textpage = page.get_textpage()
        to_px = _page_to_px(page, img.width, img.height)
        words = _pdf_words(page, textpage, to_px)
        stats = _pdf_vector_stats(page)
        rotation = page.get_rotation()
        meta: dict[str, Any] = {"rotation": rotation, "char_count": textpage.count_chars()}
        if reduced:
            meta["dpi_reduced"] = True
            reply["warnings"].append(
                f"page {i + 1}: sheet is larger than the pixel budget; rasterised at {dpi:.0f} DPI"
            )
        swap = rotation in (90, 270)
        reply["pages"].append(
            {
                "index": i,
                "kind": "pdf_page",
                "width_px": img.width,
                "height_px": img.height,
                "dpi": dpi,
                "width_mm": (h_pt if swap else w_pt) * MM_PER_PT,
                "height_mm": (w_pt if swap else h_pt) * MM_PER_PT,
                "raster": _save_png(img, out, f"p{i}.png"),
                "preview": _preview(img, out, f"p{i}_preview.png"),
                "words": _write_json(out, f"p{i}_words.json", words),
                "content": _write_json(out, f"p{i}_vector.json", stats),
                "meta": meta,
            }
        )
        textpage.close()
        page.close()
    pdf.close()
    return reply


# ---------------------------------------------------------------------------------------------
# raster images (Pillow): EXIF orientation, ICC → sRGB, multi-page TIFF, pixel limits
# ---------------------------------------------------------------------------------------------
def _to_srgb(img: Any, meta: dict[str, Any]) -> Any:
    from PIL import Image, ImageCms

    icc = img.info.get("icc_profile")
    if img.mode in ("I;16", "I;16B", "I;16L", "I"):
        arr = np.asarray(img, dtype=np.float64)
        hi = 65535.0 if arr.max() > 255 else 255.0
        img = Image.fromarray(np.clip(arr / hi * 255.0 + 0.5, 0, 255).astype(np.uint8), "L")
    if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        bg = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        img = Image.alpha_composite(bg, rgba)
        meta["alpha_flattened"] = True
    if icc:
        try:
            src = ImageCms.ImageCmsProfile(io.BytesIO(icc))
            dst = ImageCms.createProfile("sRGB")
            mode = img.mode if img.mode in ("RGB", "CMYK", "L") else "RGB"
            if img.mode != mode:
                img = img.convert(mode)
            img = ImageCms.profileToProfile(img, src, dst, outputMode="RGB")
            meta["icc_converted"] = ImageCms.getProfileDescription(src).strip()
        except (ImageCms.PyCMSError, OSError, ValueError) as e:
            meta["icc_error"] = str(e)[:200]
    return img.convert("RGB")


def _exif_meta(img: Any) -> dict[str, Any]:
    from PIL import ExifTags

    exif = img.getexif()
    if not exif:
        return {}
    out: dict[str, Any] = {}
    names = {v: k for k, v in ExifTags.TAGS.items()}
    for key in ("Make", "Model", "DateTime", "Software"):
        tag = names.get(key)
        if tag in exif:
            out[key.lower()] = str(exif[tag])[:120]
    if exif.get(names.get("GPSInfo", -1)):
        out["gps_present"] = True  # coordinates are not stored (privacy)
    return out


def task_image(req: dict[str, Any]) -> Reply:
    from PIL import Image, ImageOps, ImageSequence, UnidentifiedImageError

    src, out = Path(req["path"]), Path(req["out"])
    Image.MAX_IMAGE_PIXELS = int(req["max_pixels"])
    warnings.simplefilter("error", Image.DecompressionBombWarning)
    try:
        with Image.open(src) as probe:
            probe.verify()
        img = Image.open(src)
        frames = (
            [f.copy() for f in ImageSequence.Iterator(img)]
            if getattr(img, "n_frames", 1) > 1
            else [img]
        )
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_LIMIT_EXCEEDED,
            f"The image exceeds the pixel limit ({req['max_pixels'] / 1e6:.0f} MP).",
            "Downscale the scan to ≤ 600 DPI or split it.",
        ) from e
    except (UnidentifiedImageError, OSError, SyntaxError, ValueError) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The image cannot be decoded ({e}).",
            "Re-save it as PNG or JPEG and upload it again.",
        ) from e
    reply = _empty_reply()
    exif = _exif_meta(frames[0])
    reply["meta"] = {
        "format": img.format,
        "frames": len(frames),
        **({"exif": exif} if exif else {}),
    }
    for i, frame in enumerate(frames):
        meta: dict[str, Any] = {"mode": frame.mode}
        oriented = ImageOps.exif_transpose(frame)
        if oriented is not frame and oriented.size != frame.size:
            meta["exif_rotated"] = True
        rgb = _to_srgb(oriented if oriented is not None else frame, meta)
        dpi = frame.info.get("dpi")
        reply["pages"].append(
            {
                "index": i,
                "kind": "image",
                "width_px": rgb.width,
                "height_px": rgb.height,
                "dpi": float(dpi[0]) if dpi and dpi[0] else None,
                "raster": _save_png(rgb, out, f"p{i}.png"),
                "preview": _preview(rgb, out, f"p{i}_preview.png"),
                "meta": meta,
            }
        )
    return reply


# ---------------------------------------------------------------------------------------------
# DXF (ezdxf recover): summary + raster preview drawn with OpenCV
# ---------------------------------------------------------------------------------------------
INSUNITS_MM = {1: 25.4, 2: 304.8, 4: 1.0, 5: 10.0, 6: 1000.0, 14: 100.0}


class _Cv2Backend:
    """Minimal ezdxf drawing backend that rasterises into a numpy image (preview only)."""

    def __init__(self, bbox: Any, max_px: int) -> None:
        ext = bbox.extmax - bbox.extmin
        self.scale = (max_px - 20) / max(ext.x, ext.y, 1e-9)
        self.w = max(1, int(ext.x * self.scale) + 20)
        self.h = max(1, int(ext.y * self.scale) + 20)
        self.x0, self.y0 = bbox.extmin.x, bbox.extmin.y
        self.img = np.full((self.h, self.w, 3), 255, np.uint8)

    def _pt(self, x: float, y: float) -> tuple[int, int]:
        return round((x - self.x0) * self.scale + 10), round(
            self.h - 10 - (y - self.y0) * self.scale
        )

    def _poly(self, pts: list[Any]) -> np.ndarray:
        return np.array([self._pt(p.x, p.y) for p in pts], np.int32)

    def configure(self, config: Any) -> None:
        pass

    def set_background(self, color: Any) -> None:
        pass

    def enter_entity(self, entity: Any, properties: Any) -> None:
        pass

    def exit_entity(self, entity: Any) -> None:
        pass

    def clear(self) -> None:
        self.img[:] = 255

    def finalize(self) -> None:
        pass

    def draw_point(self, pos: Any, properties: Any) -> None:
        import cv2

        cv2.circle(self.img, self._pt(pos.x, pos.y), 1, (0, 0, 0), -1)

    def draw_line(self, start: Any, end: Any, properties: Any) -> None:
        import cv2

        cv2.line(
            self.img, self._pt(start.x, start.y), self._pt(end.x, end.y), (0, 0, 0), 1, cv2.LINE_AA
        )

    def draw_solid_lines(self, lines: Any, properties: Any) -> None:
        for a, b in lines:
            self.draw_line(a, b, properties)

    def draw_path(self, path: Any, properties: Any) -> None:
        import cv2

        for sub in path.sub_paths() if path.has_sub_paths else [path]:
            pts = list(sub.flattening(1.0 / self.scale))
            if len(pts) >= 2:
                cv2.polylines(self.img, [self._poly(pts)], False, (0, 0, 0), 1, cv2.LINE_AA)

    def draw_filled_paths(self, paths: Any, properties: Any) -> None:
        import cv2

        polys = []
        for p in paths:
            for sub in p.sub_paths() if p.has_sub_paths else [p]:
                pts = list(sub.flattening(1.0 / self.scale))
                if len(pts) >= 3:
                    polys.append(self._poly(pts))
        if polys:
            cv2.fillPoly(self.img, polys, (0, 0, 0), cv2.LINE_AA)

    def draw_filled_polygon(self, points: Any, properties: Any) -> None:
        import cv2

        pts = list(points.vertices())
        if len(pts) >= 3:
            cv2.fillPoly(self.img, [self._poly(pts)], (0, 0, 0), cv2.LINE_AA)

    def draw_image(self, image_data: Any, properties: Any) -> None:
        pass  # raster references are not embedded in DXF; they are uploaded separately


def _repair_zero_handles(src: Path, out: Path) -> tuple[Path, int]:
    """ASCII DXF: replace handle ``0`` (group codes 5/105) with fresh handles above the maximum."""
    raw = src.read_bytes()
    text = raw.decode("utf-8") if raw[:3] != b"\xef\xbb\xbf" else raw[3:].decode("utf-8")
    lines = text.splitlines()
    top = 0
    for i in range(0, len(lines) - 1, 2):
        if lines[i].strip() in ("5", "105"):
            try:
                top = max(top, int(lines[i + 1].strip(), 16))
            except ValueError:
                continue
    fixed = 0
    for i in range(0, len(lines) - 1, 2):
        if lines[i].strip() in ("5", "105") and lines[i + 1].strip() == "0":
            top += 1
            lines[i + 1] = format(top, "X")
            fixed += 1
    dest = out / "repaired.dxf"
    dest.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return dest, fixed


def task_dxf(req: dict[str, Any]) -> Reply:
    from ezdxf import recover
    from ezdxf.addons.drawing.frontend import Frontend
    from ezdxf.addons.drawing.properties import RenderContext
    from ezdxf.addons.drawing.recorder import Recorder
    from ezdxf.entities.insert import Insert
    from ezdxf.entities.mtext import MText
    from ezdxf.entities.text import Text
    from ezdxf.lldxf.const import DXFStructureError
    from PIL import Image

    src, out = Path(req["path"]), Path(req["out"])
    repaired = 0
    try:
        try:
            doc, auditor = recover.readfile(src)
        except ValueError as e:
            if "Invalid handle" not in str(e):
                raise
            # e.g. LibreDWG writes handle 0 for some ENDBLK records; give them fresh handles
            src, repaired = _repair_zero_handles(src, out)
            doc, auditor = recover.readfile(src)
    except (OSError, ValueError, DXFStructureError) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The DXF cannot be read ({e}).",
            "Re-export the drawing as DXF (R2010 or later) from the CAD application.",
        ) from e
    msp = doc.modelspace()
    layers: dict[str, dict[str, int]] = {}
    texts: list[dict[str, Any]] = []
    count = 0
    for ent in msp:
        count += 1
        t = ent.dxftype()
        layer = ent.dxf.get("layer", "0")
        layers.setdefault(layer, {})
        layers[layer][t] = layers[layer].get(t, 0) + 1
        if len(texts) >= 20_000:
            continue
        if isinstance(ent, (Text, MText)):
            txt = ent.plain_text() if isinstance(ent, MText) else ent.dxf.text
            ins = ent.dxf.insert
            height = ent.dxf.char_height if isinstance(ent, MText) else ent.dxf.height
            texts.append(
                {
                    "text": unicodedata.normalize("NFC", str(txt)),
                    "x": ins.x,
                    "y": ins.y,
                    "height": float(height or 0.0),
                    "rotation": float(ent.dxf.get("rotation", 0.0)),
                    "layer": layer,
                }
            )
        elif isinstance(ent, Insert):
            for a in ent.attribs:
                ins = a.dxf.insert
                texts.append(
                    {
                        "text": unicodedata.normalize("NFC", str(a.dxf.text)),
                        "x": ins.x,
                        "y": ins.y,
                        "height": float(a.dxf.height or 0.0),
                        "rotation": float(a.dxf.get("rotation", 0.0)),
                        "layer": a.dxf.get("layer", layer),
                        "block": ent.dxf.name,
                        "tag": a.dxf.tag,
                    }
                )
    insunits = int(doc.header.get("$INSUNITS", 0))
    summary = {
        "dxfversion": doc.dxfversion,
        "insunits": insunits,
        "unit_mm": INSUNITS_MM.get(insunits),
        "measurement": int(doc.header.get("$MEASUREMENT", 1)),
        "entity_count": count,
        "layers": layers,
        "texts": texts,
        "blocks": len(doc.blocks),
        "audit_errors": len(auditor.errors),
        "audit_fixes": len(auditor.fixes),
    }
    reply = _empty_reply()
    if auditor.has_errors:
        reply["warnings"].append(f"DXF had {len(auditor.errors)} structural errors (recovered)")
    if repaired:
        reply["warnings"].append(
            f"DXF had {repaired} records with the invalid handle 0 (renumbered)"
        )
        summary["repaired_handles"] = repaired
    page: dict[str, Any] = {"index": 0, "kind": "dxf", "meta": {"insunits": insunits}}
    if count:
        rec = Recorder()
        Frontend(RenderContext(doc), rec).draw_layout(msp, finalize=True)
        player = rec.player()
        bbox = player.bbox()
        if bbox.has_data:
            be = _Cv2Backend(bbox, int(req.get("preview_px", 4096)))
            player.replay(be)  # type: ignore[arg-type]
            img = Image.fromarray(be.img, "RGB")
            page.update(
                width_px=img.width,
                height_px=img.height,
                raster=_save_png(img, out, "p0.png"),
                preview=_preview(img, out, "p0_preview.png"),
            )
            summary["extents"] = [bbox.extmin.x, bbox.extmin.y, bbox.extmax.x, bbox.extmax.y]
            summary["raster_scale_px_per_unit"] = be.scale
    page["content"] = _write_json(out, "p0_dxf.json", summary)
    reply["pages"].append(page)
    reply["meta"] = {"dxfversion": doc.dxfversion, "insunits": insunits, "entity_count": count}
    return reply


# ---------------------------------------------------------------------------------------------
# Office (python-docx / openpyxl / python-pptx). Macros are never executed.
# ---------------------------------------------------------------------------------------------
def _export_images(parts: Iterator[Any], out: Path, reply: Reply, prefix: str) -> None:
    seen: set[str] = set()
    for part in parts:
        ct = getattr(part, "content_type", "")
        if not ct.startswith("image/") or ct in ("image/x-emf", "image/x-wmf"):
            continue
        if str(part.partname).startswith("/docProps/"):
            continue  # the file's own thumbnail, not user content
        name = PurePosixPath(str(part.partname)).name
        if name in seen:
            continue
        seen.add(name)
        dest = out / f"{prefix}_{len(seen)}_{name}"
        dest.write_bytes(part.blob)
        reply["children"].append({"path": dest.name, "name": name})


def task_docx(req: dict[str, Any]) -> Reply:
    import docx

    src, out = Path(req["path"]), Path(req["out"])
    try:
        d = docx.Document(str(src))
    except (ValueError, KeyError, zipfile.BadZipFile) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The Word document cannot be read ({e}).",
            "Save it as a plain .docx (not .docm) and upload it again.",
        ) from e
    paragraphs = [
        {
            "text": unicodedata.normalize("NFC", p.text),
            "style": p.style.name if p.style is not None else "",
        }
        for p in d.paragraphs
        if p.text.strip()
    ]
    tables = [
        [[unicodedata.normalize("NFC", c.text.strip()) for c in row.cells] for row in t.rows]
        for t in d.tables
    ]
    reply = _empty_reply()
    reply["pages"].append(
        {
            "index": 0,
            "kind": "docx",
            "content": _write_json(
                out, "p0_docx.json", {"paragraphs": paragraphs, "tables": tables}
            ),
            "meta": {"paragraphs": len(paragraphs), "tables": len(tables)},
        }
    )
    _export_images(iter(d.part.package.iter_parts()), out, reply, "img")
    return reply


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, (int, float, bool)):
        return v
    return unicodedata.normalize("NFC", str(v))


def task_xlsx(req: dict[str, Any]) -> Reply:
    import openpyxl

    src, out = Path(req["path"]), Path(req["out"])
    # a file object: openpyxl rejects paths without an Excel extension (uploads are staged as .bin);
    # it stays open for the read-only workbook until the task process exits
    fh = src.open("rb")
    try:
        wb = openpyxl.load_workbook(
            fh, read_only=True, data_only=True, keep_vba=False, keep_links=False
        )
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The workbook cannot be read ({e}).",
            "Save it again as .xlsx and upload it.",
        ) from e
    reply = _empty_reply()
    budget = int(req["max_cells"])
    for i, ws in enumerate(wb.worksheets):
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            budget -= len(row)
            if budget < 0:
                raise ArchRenderError(
                    ErrorCode.INGEST_LIMIT_EXCEEDED,
                    f"The workbook has more than {req['max_cells']} cells.",
                    "Remove unrelated sheets or split the workbook.",
                )
            rows.append([_cell(v) for v in row])
        while rows and all(v in (None, "") for v in rows[-1]):
            rows.pop()
        reply["pages"].append(
            {
                "index": i,
                "kind": "sheet",
                "content": _write_json(out, f"p{i}_sheet.json", {"name": ws.title, "rows": rows}),
                "meta": {"sheet": ws.title, "rows": len(rows)},
            }
        )
    reply["meta"] = {"sheets": len(wb.worksheets), "vba_ignored": bool(req.get("macro"))}
    wb.close()
    return reply


def task_pptx(req: dict[str, Any]) -> Reply:
    import pptx

    src, out = Path(req["path"]), Path(req["out"])
    try:
        prs = pptx.Presentation(str(src))
    except (ValueError, KeyError, zipfile.BadZipFile) as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"The presentation cannot be read ({e}).",
            "Save it as a plain .pptx (not .pptm) and upload it again.",
        ) from e
    reply = _empty_reply()
    for i, slide in enumerate(prs.slides):
        texts, tables, pictures = [], [], 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(unicodedata.normalize("NFC", shape.text_frame.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append([[c.text for c in r.cells] for r in shape.table.rows])
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                pictures += 1
        reply["pages"].append(
            {
                "index": i,
                "kind": "slide",
                "content": _write_json(out, f"p{i}_slide.json", {"texts": texts, "tables": tables}),
                "meta": {"pictures": pictures},
            }
        )
    _export_images(iter(prs.part.package.iter_parts()), out, reply, "img")
    return reply


# ---------------------------------------------------------------------------------------------
# ZIP: safe extraction under limits (bombs, traversal, links, devices, encryption)
# ---------------------------------------------------------------------------------------------
_JUNK = ("__MACOSX/", ".DS_Store", "Thumbs.db", "desktop.ini")


def _unsafe(
    msg: str, hint: str = "Re-create the archive from ordinary files and folders."
) -> ArchRenderError:
    return ArchRenderError(ErrorCode.INGEST_UNSAFE_ARCHIVE, msg, hint)


def _safe_name(raw: str) -> str:
    name = unicodedata.normalize("NFC", raw.replace("\\", "/"))
    if name.startswith("/") or (len(name) > 1 and name[1] == ":"):
        raise _unsafe(f"Archive entry {raw!r} has an absolute path.")
    parts = [p for p in name.split("/") if p not in ("", ".")]
    if any(p == ".." for p in parts):
        raise _unsafe(f"Archive entry {raw!r} points outside the archive ('..').")
    if any("\x00" in p for p in parts):
        raise _unsafe(f"Archive entry {raw!r} has an invalid name.")
    return "/".join(parts)


def task_unzip(req: dict[str, Any]) -> Reply:
    src, out = Path(req["path"]), Path(req["out"])
    max_total, max_ratio = int(req["max_total_bytes"]), float(req["max_ratio"])
    max_entries = int(req["max_entries"])
    try:
        zf = zipfile.ZipFile(src)
    except zipfile.BadZipFile as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT, f"Corrupt ZIP archive ({e}).", "Re-create the ZIP."
        ) from e
    infos = zf.infolist()
    if len(infos) > max_entries:
        raise _unsafe(
            f"The archive has {len(infos)} entries; the limit is {max_entries}.",
            "Split the archive.",
        )
    declared = sum(i.file_size for i in infos)
    compressed = sum(i.compress_size for i in infos)
    if declared > max_total:
        raise ArchRenderError(
            ErrorCode.INGEST_LIMIT_EXCEEDED,
            f"The archive expands to {declared / 1e9:.1f} GB; the limit is {max_total / 1e9:.1f} GB.",
            "Split the archive.",
        )
    if compressed and declared / compressed > max_ratio and declared > 10 * 1024 * 1024:
        raise _unsafe(
            f"The archive compression ratio is {declared / compressed:.0f}:1 (zip bomb?)."
        )
    reply = _empty_reply()
    names: set[str] = set()
    written = 0
    for info in infos:
        mode = (info.external_attr >> 16) & 0xFFFF
        if stat.S_ISLNK(mode):
            raise _unsafe(f"Archive entry {info.filename!r} is a symbolic link.")
        if mode and (
            stat.S_ISCHR(mode) or stat.S_ISBLK(mode) or stat.S_ISFIFO(mode) or stat.S_ISSOCK(mode)
        ):
            raise _unsafe(f"Archive entry {info.filename!r} is a device or special file.")
        name = _safe_name(info.filename)
        if info.is_dir() or not name:
            continue
        if any(j in info.filename for j in _JUNK):
            continue
        if info.flag_bits & 0x1:
            raise ArchRenderError(
                ErrorCode.INGEST_ENCRYPTED,
                f"Archive entry {info.filename!r} is encrypted.",
                "Create the ZIP without a password.",
            )
        if (
            info.compress_size
            and info.file_size / info.compress_size > max_ratio
            and info.file_size > 1024 * 1024
        ):
            raise _unsafe(
                f"Archive entry {info.filename!r} has a compression ratio above {max_ratio:.0f}:1."
            )
        key = name.casefold()
        if key in names:
            reply["warnings"].append(f"duplicate entry {name!r} skipped")
            continue
        names.add(key)
        dest = out / "files" / f"{len(names):05d}"
        dest.parent.mkdir(parents=True, exist_ok=True)
        size = 0
        with zf.open(info) as fh, dest.open("wb") as w:
            while buf := fh.read(1 << 20):
                size += len(buf)
                written += len(buf)
                if size > info.file_size or written > max_total:
                    raise _unsafe(
                        f"Archive entry {info.filename!r} expands beyond its declared size."
                    )
                w.write(buf)
        reply["children"].append({"path": str(dest.relative_to(out)), "name": name})
    reply["meta"] = {"entries": len(reply["children"]), "uncompressed_bytes": written}
    return reply


def task_zip_names(req: dict[str, Any]) -> Reply:
    """Entry names of a ZIP container (for OOXML type detection), without extracting anything."""
    try:
        with zipfile.ZipFile(Path(req["path"])) as zf:
            names = zf.namelist()[: int(req["max_entries"])]
    except zipfile.BadZipFile as e:
        raise ArchRenderError(
            ErrorCode.INGEST_CORRUPT,
            f"Corrupt ZIP container ({e}).",
            "Re-export or re-zip the file.",
        ) from e
    reply = _empty_reply()
    reply["meta"] = {"names": names}
    return reply


TASKS: dict[str, Callable[[dict[str, Any]], Reply]] = {
    "pdf": task_pdf,
    "image": task_image,
    "dxf": task_dxf,
    "docx": task_docx,
    "xlsx": task_xlsx,
    "pptx": task_pptx,
    "unzip": task_unzip,
    "zip_names": task_zip_names,
}


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 1 or argv[0] not in TASKS:
        print(json.dumps({"ok": False, "code": "INTERNAL", "message": f"unknown task {argv}"}))
        return 2
    req = json.loads(sys.stdin.read())
    Path(req["out"]).mkdir(parents=True, exist_ok=True)
    try:
        result = TASKS[argv[0]](req)
        print(json.dumps({"ok": True, "result": result}, ensure_ascii=False))
        return 0
    except ArchRenderError as e:
        print(
            json.dumps(
                {"ok": False, "code": e.code.value, "message": e.message, "fix_hint": e.fix_hint}
            )
        )
        return 1
    except MemoryError:
        print(
            json.dumps(
                {
                    "ok": False,
                    "code": ErrorCode.INGEST_LIMIT_EXCEEDED.value,
                    "message": "The parser ran out of its memory allowance.",
                    "fix_hint": "Split the file or ask an admin to raise ARCHRENDER_SANDBOX_MEMORY_MB.",
                }
            )
        )
        return 1
    except Exception as e:
        print(
            json.dumps(
                {
                    "ok": False,
                    "code": ErrorCode.INGEST_CORRUPT.value,
                    "message": f"The file could not be parsed ({type(e).__name__}: {str(e)[:300]}).",
                    "fix_hint": "Re-export the file from the authoring application and upload it again.",
                    "context": {"traceback": traceback.format_exc()[-3000:]},
                }
            )
        )
        return 1


if __name__ == "__main__":
    sys.exit(main())
