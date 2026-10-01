"""Golden projects (docs/PLAN.md): the document sets S0–S2 must handle, with ground truth.

- **G1 "Daire"**: Turkish 3+1-style flat. Vector plan PDF (room numbers + door/window tags), the same
  plan as DXF (DUVAR/KAPI/PENCERE, centimetres), finish schedule XLSX (Mahal Listesi), a
  door/window schedule PDF, a Turkish brief (DOCX) and mood-board images.
- **G2 "Loft"**: a skewed plan with an arc wall and a double-height living room: the plan as a noisy
  300 DPI scan with an approval stamp, a phone photo of the printed sheet, the reflected ceiling
  plan (PDF, ceiling heights), the door/window schedule PDF and a Turkish/English DOCX brief.
- **G3 "Office"**: an IFC model (two storeys) and the ground floor as an English PDF in feet and
  inches (1/4" = 1'-0").

``spec`` is the ground-truth plan; each plan document carries the transform from plan metres into
its own frame (``GoldenDoc.gt``: ``plan_to_page_px``, ``plan_to_doc`` or ``plan_to_photo_h``).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field, replace
from typing import Any

import docx
import numpy as np
import openpyxl

from archrender.synth.dxf import plan_dxf
from archrender.synth.ifc import IfcStorey, plan_ifc
from archrender.synth.layout import Layout, random_layout
from archrender.synth.plan import PlanSpec, add_loft_features, from_layout
from archrender.synth.raster import phone_photo, photo, scan
from archrender.synth.sheets import (
    BRIEF_EN,
    BRIEF_TR,
    MATS_TR,
    floor_plan_page,
    moodboard_page,
    schedule_page,
    schedule_rows,
)
from archrender.understand.text import normalise_tag


@dataclass
class GoldenDoc:
    filename: str
    data: bytes
    expected_class: str | None  # None: container/derived, not classified as one page
    gt: dict[str, Any] | None = None  # plan → document frame for plan documents


@dataclass
class GoldenProject:
    name: str
    layout: Layout
    docs: list[GoldenDoc] = field(default_factory=list)
    schedule_tags: set[str] = field(default_factory=set)  # normalised tags every row should link to
    spec: PlanSpec | None = None  # ground-truth plan (the ground floor for G3)
    storeys: list[PlanSpec] = field(default_factory=list)  # G3: every storey


def _xlsx(title: str, header: list[str], rows: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:31]
    ws.append([title])
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _brief_docx(rng: np.random.Generator, bilingual: bool) -> bytes:
    d = docx.Document()
    d.add_heading("Tasarım Özeti", level=1)
    sentences = list(BRIEF_TR)
    rng.shuffle(sentences)
    for s in sentences[:5]:
        d.add_paragraph(s.format(mat=rng.choice(MATS_TR)))
    if bilingual:
        d.add_heading("Design brief (English)", level=1)
        for s in BRIEF_EN[:4]:
            d.add_paragraph(s.format(mat="oak"))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def g1_daire(seed: int = 101) -> GoldenProject:
    rng = np.random.default_rng(seed)
    layout = random_layout(rng)
    spec = from_layout(layout)
    plan = floor_plan_page(rng, layout, room_numbers=True)
    dxf, dxf_gt = plan_dxf(spec, unit="cm")
    dw = schedule_page(rng, layout, kind="door_window")
    header, rows = schedule_rows(layout, "finish", "tr", rng)
    g = GoldenProject("G1 Daire", layout, spec=spec)
    g.docs = [
        GoldenDoc("Zemin Kat Planı.pdf", plan.pdf, "floor_plan", plan.gt),
        GoldenDoc("Zemin Kat Planı.dxf", dxf, "floor_plan", dxf_gt),
        GoldenDoc("Malzeme Listesi.xlsx", _xlsx("Mahal Listesi", header, rows), "schedule"),
        GoldenDoc("Doğrama Listesi.pdf", dw.pdf, "schedule"),
        GoldenDoc("Tasarım Özeti.docx", _brief_docx(rng, bilingual=False), None),
        GoldenDoc("Konsept Panosu.pdf", moodboard_page(rng).pdf, "moodboard"),
    ]
    for i in range(2):
        data, _ = photo(rng)
        g.docs.append(GoldenDoc(f"Referans {i + 1}.jpg", data, "photo"))
    g.schedule_tags = {
        normalise_tag(t) or t for t in [o.tag for o in layout.openings] + [r[0] for r in rows]
    }
    return g


def g2_loft(seed: int = 202) -> GoldenProject:
    rng = np.random.default_rng(seed)
    layout = random_layout(rng)
    spec = from_layout(layout, variant="skewed_arc", rng=rng)
    add_loft_features(spec)
    plan = floor_plan_page(rng, spec=spec, wall_style="solid", stamp=True)
    scan_bytes, _, scan_gt = scan(plan.pdf, plan.gt, rng, dpi=300, quality="noisy")
    photo_bytes, photo_gt = phone_photo(plan.pdf, plan.gt, rng)
    rcp = floor_plan_page(rng, spec=spec, ceiling=True)
    # the schedule lists the openings the plan draws (the arc wall has no windows)
    drawn = {o.tag for o in spec.openings}
    dw = schedule_page(
        rng,
        replace(layout, openings=[o for o in layout.openings if o.tag in drawn]),
        kind="door_window",
    )
    g = GoldenProject("G2 Loft", layout, spec=spec)
    g.docs = [
        GoldenDoc("Zemin Kat Planı (tarama).jpg", scan_bytes, "floor_plan", scan_gt),
        GoldenDoc("Zemin Kat Planı (telefon).jpg", photo_bytes, "floor_plan", photo_gt),
        GoldenDoc("Tavan Planı.pdf", rcp.pdf, "ceiling_plan", rcp.gt),
        GoldenDoc("Doğrama Listesi.pdf", dw.pdf, "schedule"),
        GoldenDoc("Brief TR-EN.docx", _brief_docx(rng, bilingual=True), None),
    ]
    g.schedule_tags = drawn
    return g


def g3_office(seed: int = 303) -> GoldenProject:
    rng = np.random.default_rng(seed)
    layout = random_layout(rng, english=True)
    ground = from_layout(layout, level_name="Ground Floor", rng=rng)
    upper = from_layout(layout, level_name="Level 1", rng=rng)
    plan = floor_plan_page(rng, spec=ground, lang="en", imperial=True, wall_style="solid")
    g = GoldenProject("G3 Office", layout, spec=ground, storeys=[ground, upper])
    g.docs = [
        GoldenDoc(
            "Office.ifc",
            plan_ifc([IfcStorey(ground, "Ground Floor", 0.0), IfcStorey(upper, "Level 1", 3.2)]),
            None,
            {"plan_to_doc": [[1.0, 0.0, 0.0], [0.0, 1.0, 0.0]]},
        ),
        GoldenDoc("Ground Floor Plan.pdf", plan.pdf, "floor_plan", plan.gt),
    ]
    return g
